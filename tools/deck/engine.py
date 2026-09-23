"""좌표 하나로 pptx 와 HTML 미리보기를 동시에 뽑는다.

**왜 이렇게 하나.** 이 기계에 LibreOffice 가 없어 pptx 를 그림으로 못 뽑는다.
그러면 배치를 새로 짜는 순간 눈으로 확인할 방법이 사라진다. 그래서 같은 좌표
목록에서 HTML 도 함께 만들고, 그것을 Chrome 으로 찍어 검사한다 — 글꼴도 좌표도
같으므로 미리보기가 거짓말을 할 여지가 좁다.

단위는 pt 다. 캔버스 1440x810pt = 20x11.25in, 3주차 덱과 같다.
"""
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import html as _html

W, H = 1440.0, 810.0
EMU = 12700

# ── 3주차 덱에서 그대로 뜬 값들 ────────────────────────────────────────
INK    = "111737"   # 본문 잉크 · 어두운 띠 배경
NAVY   = "1E2761"   # 강조 라벨 · 번호 배지
MUTED  = "515B73"   # 보조 텍스트
ICE    = "CADCFC"   # 어두운 띠 위 본문
PANEL  = "F2F6FC"   # 밝은 패널
LINE   = "D6DEEF"   # 얇은 구분선
TRACK  = "DCE3F2"   # 진행바 바닥
GREEN  = "2E7D5B"
RED    = "E8543F"
WHITE  = "FFFFFF"
RULE   = "3A4474"

FONT, FONTB = "IBM Plex Sans KR", "IBM Plex Sans KR Bold"

M   = 75.6          # 좌우 여백
CW  = 1285.2        # 본문 폭
PAD = 28.1          # 패널 안쪽 여백

# ── 세로 자리 — **한 곳에서만 지킨다** ────────────────────────────────
#
# **한때 장마다 손으로 박혀 있었다.** 띠가 464 · 470 · 500 셋이었고 꼬리가
# 그때마다 «띠 + 158» 로 따라 적혔다 — **여덟 곳에 같은 관계가 손으로** 있었다.
# 띠 높이를 한 번 바꾸면 여덟 곳을 다 고쳐야 하고, 하나 빠뜨리면 그 장만
# 어긋난다(#217).
#
# **#216 이 자료 쪽에서 이미 6.94" 로 통일했다.** 그 값이 여기 BAND_Y 다 —
# 셋 중 가장 아래이므로 어느 장에서도 내용과 겹치지 않는다.
#
# 인치는 자료를 실측한 값이다(#216 「레이아웃 정합」). pt = 인치 × 72.
BAND_Y       = 500.0    # 6.94" — 어두운 띠 위쪽
BAND_H       = 133.9    # 바닥이 8.80" 가 된다. **#216 이 렌더링으로 확인한 값**
BAND_HEAD_DY = 23.8     # 띠 안 제목
BAND_BODY_DY = 61.9     # 0.86" — 띠 안 본문
NOTE_DY      = 155.2    # 꼬리는 띠에서 이만큼 아래. BAND_Y + NOTE_DY = 9.10"
SUB_Y        = 120.96   # 1.68" — 머리글 부제

# **띠 본문은 2줄이 상한이다**(#216). 어두운 띠가 바닥 8.80" 에 고정이라 셋째
# 줄이 잘린다 — 초안이 3줄이 되어 실제로 잘렸고, 줄여서 내보냈다. 문서에만
# 적으면 다음에 또 밟으므로 생성기가 센다.
BAND_BODY_MAX_LINES = 2
BAND_BODY_CPL = 62      # 19.5pt 로 1188pt 폭에 들어가는 대략의 글자 수


class Slide:
    def __init__(self, bg=WHITE):
        self.bg = bg
        self.ops = []
        self.speaker_note = ""

    def note_text(self, t):
        """발표자 노트. **생성기가 노트를 안 만들면 재생성 때마다 노트가
        사라진다** — #217·#224 가 그 사고의 기록이다."""
        self.speaker_note = t
        return self

    # ── 원시 ───────────────────────────────────────────────────────────
    def rect(self, x, y, w, h, fill, r=8.0):
        self.ops.append(("rect", x, y, w, h, fill, r))
        return self

    def text(self, x, y, w, h, lines, size, color=INK, bold=False,
             align="l", anchor="t", lh=1.2):
        if isinstance(lines, str):
            lines = [lines]
        self.ops.append(("text", x, y, w, h, list(lines), size, color,
                         bold, align, anchor, lh))
        return self

    def image(self, path, x=0, y=0, w=W, h=H):
        """그림. **없으면 그 자리를 비우고 경고한다 — 죽지 않는다.**

        그림은 빌드 산출물이다(SVG 정본에서 뽑는다). 없다고 나머지 열 장까지
        못 뽑을 이유가 없고, 그렇다고 조용히 빠지면 «빈 장으로 발표장에 간다».
        노트와 같은 처리다."""
        import warnings
        from pathlib import Path
        if not Path(path).exists():
            warnings.warn(f"그림이 없다: {path} — **그 자리를 비우고 뽑는다.** "
                          f"뽑는 법은 tools/deck/README.md 에 있다.", stacklevel=2)
            return self
        self.ops.append(("image", x, y, w, h, path))
        return self

    # ── 조립 ───────────────────────────────────────────────────────────
    def header(self, no, title, sub):
        self.rect(M, 56.2, 49.7, 49.7, NAVY, r=10)
        self.text(M, 56.2, 49.7, 49.7, no, 21, WHITE, True, "c", "m")
        self.text(142.6, 22.9, 1026, 116.1, title, 49.5, INK, True, "l", "m")
        self.text(142.6, SUB_Y, 1231.2, 27, sub, 22.5, MUTED)
        return self

    def band(self, heading, body, y=BAND_Y, h=BAND_H):
        """어두운 띠. **y 를 안 주면 통일된 자리다** — 그게 기본이어야 한다.

        y 를 인자로 남겨 두는 이유는 예외를 허용하려는 것이 아니라, 예외가
        필요해졌을 때 **그 장에서만 눈에 띄게** 하려는 것이다.
        """
        import warnings
        n = _band_body_lines(body)
        if n > BAND_BODY_MAX_LINES:
            warnings.warn(
                f"띠 본문이 {n}줄로 보인다 — 상한은 {BAND_BODY_MAX_LINES}줄이다. "
                f"띠 바닥이 8.80\" 에 고정이라 셋째 줄은 잘린다(#216). "
                f"«{heading}»", stacklevel=2)
        self.rect(M, y, CW, h, INK, r=10)
        self.text(124.2, y + BAND_HEAD_DY, 1188, 30, heading, 25.5, WHITE, True)
        self.text(124.2, y + BAND_BODY_DY, 1188, h - 78, body, 19.5, ICE, lh=1.5)
        return self

    def note(self, t, y=BAND_Y + NOTE_DY):
        """띠 아래 꼬리. **기본값이 띠에서 파생된다** — 둘의 간격을 두 곳에서
        지키지 않는다."""
        return self.text(M, y, CW, 26.2, t, 18.75, MUTED)

    def panels(self, y, h, n, gap=24.0):
        """n 등분 패널을 깔고 (x, w) 목록을 돌려준다."""
        w = (CW - gap * (n - 1)) / n
        out = []
        for i in range(n):
            x = M + i * (w + gap)
            self.rect(x, y, w, h, PANEL)
            out.append((x, w))
        return out


class Deck:
    def __init__(self):
        self.slides = []

    def add(self, bg=WHITE):
        s = Slide(bg)
        self.slides.append(s)
        return s

    # ── 발표자 노트 ────────────────────────────────────────────────────
    def load_notes(self, path):
        """노트를 **저장소 밖 파일에서** 읽어 장 번호대로 붙인다(#217 결정 1).

        **왜 파일이 따로인가.** 노트에는 「안 할 말」과 예상 질문이 들어가는데
        이 저장소는 공개다(#227). 그렇다고 `.pptx` 안에만 두면 **다시 뽑을
        때마다 사라진다** — 실제로 11장 중 10장을 한 번 잃었다(#217).

        그래서 생성기는 저장소에, 노트는 `docs/submissions/` 아래에 둔다.
        **거기는 `docs/submissions/*` 가 이미 무시한다**(.gitignore). 그러면
        사본이 둘이 아니라 하나가 되고 `.pptx` 는 그 파일의 **산출물**이 된다.

        **없으면 멈추지 않고 경고만 한다.** 배치를 고치는 사람과 노트를 쓰는
        사람이 같지 않을 수 있고, 노트가 없다고 그림까지 못 뽑을 이유는 없다.
        대신 «노트 없이 뽑았다»가 조용히 지나가면 안 되므로 소리를 낸다.

        형식은 `### 슬라이드 N` 이고, 그 아래가 그 장의 노트다.
        """
        import re, warnings
        from pathlib import Path
        p = Path(path)
        if not p.exists():
            warnings.warn(
                f"노트 파일이 없다: {p} — **노트 없이 뽑는다.** "
                f"쓰는 자리는 docs/submissions/ 아래이고 저장소에 안 올라간다(#217).",
                stacklevel=2)
            return self

        body = p.read_text(encoding="utf-8")
        found, n = {}, None
        for line in body.splitlines():
            m = re.match(r"^###\s*슬라이드\s*(\d+)", line.strip())
            if m:
                n = int(m.group(1))
                found[n] = []
            elif n is not None:
                found[n].append(line)

        for n, lines in found.items():
            if 1 <= n <= len(self.slides):
                self.slides[n - 1].note_text("\n".join(lines).strip())

        missing = [i + 1 for i, s in enumerate(self.slides) if not s.speaker_note]
        extra = [n for n in found if n > len(self.slides)]
        if missing:
            warnings.warn(f"노트가 없는 장: {missing}", stacklevel=2)
        if extra:
            warnings.warn(f"장 수보다 큰 번호가 노트에 있다: {extra}", stacklevel=2)
        print(f"[deck] 노트 {len(found) - len(extra)}장을 {p.name} 에서 읽었다.")
        return self

    # ── pptx ───────────────────────────────────────────────────────────
    def save_pptx(self, path):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(int(W * EMU)), Emu(int(H * EMU))
        blank = prs.slide_layouts[6]
        for sl in self.slides:
            s = prs.slides.add_slide(blank)
            if sl.bg != WHITE:
                bgr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                         Emu(int(W * EMU)), Emu(int(H * EMU)))
                _style(bgr, sl.bg)
            if sl.speaker_note:
                s.notes_slide.notes_text_frame.text = sl.speaker_note
            for op in sl.ops:
                if op[0] == "rect":
                    _, x, y, w, h, fill, r = op
                    shp = s.shapes.add_shape(
                        MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,
                        Emu(int(x * EMU)), Emu(int(y * EMU)),
                        Emu(int(w * EMU)), Emu(int(h * EMU)))
                    _style(shp, fill)
                    if r:
                        shp.adjustments[0] = min(0.5, r / min(w, h))
                elif op[0] == "image":
                    _, x, y, w, h, p = op
                    s.shapes.add_picture(str(p), Emu(int(x * EMU)), Emu(int(y * EMU)),
                                         Emu(int(w * EMU)), Emu(int(h * EMU)))
                else:
                    _, x, y, w, h, lines, size, color, bold, align, anchor, lh = op
                    tb = s.shapes.add_textbox(
                        Emu(int(x * EMU)), Emu(int(y * EMU)),
                        Emu(int(w * EMU)), Emu(int(h * EMU)))
                    tf = tb.text_frame
                    tf.word_wrap = True
                    tf.margin_left = tf.margin_right = 0
                    tf.margin_top = tf.margin_bottom = 0
                    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP,
                                          "m": MSO_ANCHOR.MIDDLE}[anchor]
                    for i, ln in enumerate(lines):
                        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                        p.alignment = {"l": PP_ALIGN.LEFT,
                                       "c": PP_ALIGN.CENTER}[align]
                        _lnspc(p, size * lh)
                        r_ = p.add_run()
                        r_.text = ln
                        f = r_.font
                        f.name = FONTB if bold else FONT
                        f.size = Pt(size)
                        f.bold = bold
                        f.color.rgb = RGBColor.from_string(color)
        prs.save(path)

    # ── HTML 미리보기 ──────────────────────────────────────────────────
    def save_html(self, outdir):
        from pathlib import Path
        outdir = Path(outdir)
        outdir.mkdir(parents=True, exist_ok=True)
        for n, sl in enumerate(self.slides, 1):
            parts = []
            for op in sl.ops:
                if op[0] == "rect":
                    _, x, y, w, h, fill, r = op
                    parts.append(
                        f'<div style="position:absolute;left:{x}px;top:{y}px;'
                        f'width:{w}px;height:{h}px;background:#{fill};'
                        f'border-radius:{r}px"></div>')
                elif op[0] == "image":
                    _, x, y, w, h, p = op
                    parts.append(
                        f'<img src="{Path(p).name}" style="position:absolute;'
                        f'left:{x}px;top:{y}px;width:{w}px;height:{h}px">')
                else:
                    _, x, y, w, h, lines, size, color, bold, align, anchor, lh = op
                    # pre-wrap 이 있어야 «·» 앞뒤의 연속 공백이 pptx 와 같이 보인다.
                    inner = "".join(
                        f'<div style="line-height:{size*lh}px;white-space:pre-wrap">'
                        f'{_html.escape(l) or "&nbsp;"}</div>' for l in lines)
                    just = {"t": "flex-start", "m": "center"}[anchor]
                    parts.append(
                        f'<div style="position:absolute;left:{x}px;top:{y}px;'
                        f'width:{w}px;height:{h}px;display:flex;flex-direction:column;'
                        f'justify-content:{just};font-size:{size}px;color:#{color};'
                        f'font-weight:{700 if bold else 400};'
                        f'text-align:{"center" if align=="c" else "left"};'
                        f'outline:1px dashed rgba(255,0,0,.28)">{inner}</div>')
            (outdir / f"s{n:02}.html").write_text(_PAGE.format(
                bg=sl.bg, body="".join(parts)), encoding="utf-8")


_PAGE = """<meta charset="utf-8">
<link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans+KR:wght@400;500;600;700&display=swap" rel="stylesheet">
<style>
  html,body{{margin:0;padding:0;background:#fff}}
  body{{width:1440px;height:810px;position:relative;overflow:hidden;
       background:#{bg};font-family:'IBM Plex Sans KR',sans-serif;
       word-break:keep-all;-webkit-font-smoothing:antialiased}}
  div{{box-sizing:border-box}}
</style>
{body}"""


def _band_body_lines(body):
    """띠 본문이 몇 줄로 보일지 어림한다. **정확할 필요는 없다** — 2줄 상한을
    넘길 만큼 긴지를 가리면 되고, 애매하면 사람이 렌더링으로 본다(#216)."""
    if isinstance(body, str):
        body = [body]
    return sum(max(1, -(-len(t) // BAND_BODY_CPL)) for t in body)


def _style(shp, fill):
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False


def _lnspc(p, pts):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:lnSpc",):
        old = pPr.find(qn(tag))
        if old is not None:
            pPr.remove(old)
    ln = pPr.makeelement(qn("a:lnSpc"), {})
    pts_el = ln.makeelement(qn("a:spcPts"), {"val": str(int(round(pts * 100)))})
    ln.append(pts_el)
    pPr.insert(0, ln)


from pathlib import Path  # noqa: E402  (HTML 쪽에서 쓴다)
